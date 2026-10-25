"""Tests for the released EditableFigureBench freeze, runner, and scorer."""

import csv
import hashlib
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from rfs.evaluation import editable_figure_bench as bench
from rfs.evaluation import editable_figure_domains as domains
from rfs.evaluation import editable_figure_matching as matching
from rfs.evaluation.native_editability import measure_native_editability


def _write_jsonl(path: Path, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n", encoding="utf-8")


def _write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _png(path: Path, colour: str = "white") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (12, 12), colour).save(path)


def _deck_with_connector(path: Path, picture: Path, text: str = "Alpha module") -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2.0), Inches(0.5)).text_frame.text = text
    slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(1.5), Inches(0.8))
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(2.5), Inches(3.0), Inches(2.5))
    slide.shapes.add_picture(str(picture), Inches(3.5), Inches(0.5), Inches(1.0), Inches(1.0))
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))


GOOD = bench.GOOD_LABELS[0]
GOOD_LIVELY = bench.GOOD_LABELS[1]
BAD = bench.BAD_LABEL


class FigureBenchFixture(unittest.TestCase):
    """Builds a self-contained synthetic corpus so tests never read frozen data."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self._build_corpus()

    def tearDown(self):
        self._tmp.cleanup()

    def _paper(self, paper_id: str, split: str, label: str) -> dict:
        return {
            "paper_id": paper_id,
            "arxiv_id": paper_id,
            "title": "Paper " + paper_id,
            "abstract": "About " + paper_id,
            "paper_path": "experiments/literature_review/papers/pdf/" + paper_id + ".pdf",
            "split": split,
            "final_label": label,
        }

    def _figure(self, paper_id: str, source: str, rank: int, split: str, label: str) -> dict:
        folder = "_reselect_good" if source == "reselection_good" else "_reselect" if source == "reselection" else "_stage"
        return {
            "figure_id": paper_id + "__" + source + "_" + str(rank),
            "paper_id": paper_id,
            "arxiv_id": paper_id,
            "figure_path": "experiments/literature_review/papers/" + folder + "/" + paper_id + "/cand_" + str(rank) + ".png",
            "candidate_rank": rank,
            "candidate_source": source,
            "split": split,
            "final_label": label,
            "image_sha256": "sha",
            "scope": "unknown",
            "vlm_choice": "A",
        }

    def _contract(self, paper_id: str, entities, relations) -> dict:
        return {
            "paper_id": paper_id,
            "schema_version": "judge-paper-evidence-v001",
            "entities": [{"entity_id": name, "description": name} for name in entities],
            "relations": [],
            "expected_figure_content": {
                "required_entities": list(entities),
                "required_relations": list(relations),
                "optional_entities": [],
                "must_not_invent": ["centralized_execution_module"],
            },
            "evidence_spans": [{"evidence_id": "evidence_1", "source_id": "p1_c1", "page": 1, "quote": "quote"}],
            "extraction": {"model": "test-model", "prompt_version": "prompt-v001"},
        }

    def _categories_csv(self, rows=None) -> None:
        path = self.root / "experiments/literature_review/papers/papers.csv"
        path.parent.mkdir(parents=True, exist_ok=True)
        rows = rows if rows is not None else [
            {"arxiv_id": "p1", "primary_category": "cs.CL", "categories": "cs.CL; cs.AI"},
            {"arxiv_id": "p2", "primary_category": "cs.AI", "categories": "cs.AI; cs.MA"},
            {"arxiv_id": "p3", "primary_category": "cs.CR", "categories": "cs.CR"},
            {"arxiv_id": "p4", "primary_category": "cs.CV", "categories": "cs.CV"},
            {"arxiv_id": "p5", "primary_category": "cs.LG", "categories": "cs.LG"},
            {"arxiv_id": "p6", "primary_category": "cs.AI", "categories": "cs.AI"},
            {"arxiv_id": "p7", "primary_category": "cs.AI", "categories": "cs.AI; cs.CL"},
            {"arxiv_id": "p8", "primary_category": "cs.MA", "categories": "cs.MA; cs.LG"},
            {"arxiv_id": "p9", "primary_category": "cs.IR", "categories": "cs.IR"},
        ]
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=["arxiv_id", "primary_category", "categories"])
            writer.writeheader()
            writer.writerows(rows)

    def _build_corpus(self):
        papers = [
            self._paper("p1", "test", GOOD),
            self._paper("p2", "validation", GOOD_LIVELY),
            self._paper("p3", "test", BAD),
            self._paper("p4", "validation", BAD),
            self._paper("p5", "train", GOOD),
            self._paper("p6", "test", GOOD),
            self._paper("p7", "test", GOOD),
            self._paper("p8", "test", GOOD),
            self._paper("p9", "validation", GOOD),
        ]
        figures = [
            self._figure("p1", "initial", 1, "test", GOOD),
            self._figure("p2", "initial", 1, "validation", GOOD_LIVELY),
            self._figure("p3", "initial", 1, "test", BAD),
            self._figure("p4", "initial", 1, "validation", BAD),
            self._figure("p5", "initial", 1, "train", GOOD),
            self._figure("p6", "initial", 1, "test", GOOD),
            self._figure("p7", "initial", 1, "test", GOOD),
            self._figure("p8", "initial", 1, "test", GOOD),
            self._figure("p9", "initial", 1, "validation", GOOD),
            self._figure("p9", "reselection_good", 1, "validation", GOOD),
        ]
        reviews = [
            {"arxiv_id": paper_id, "decision": "selected", "selected_ranks": [1], "selected_candidate_source": "initial", "notes": ""}
            for paper_id in ("p1", "p2", "p3", "p4", "p5", "p7", "p8", "p9")
        ]
        reviews.append({"arxiv_id": "p6", "decision": "selected", "selected_ranks": [1], "selected_candidate_source": "initial", "notes": "figure 2 is the wrong figure"})

        _write_jsonl(self.root / "datasets/judge-framework-v001/papers.jsonl", papers)
        _write_jsonl(self.root / "datasets/judge-framework-v001/figures.jsonl", figures)
        _write_jsonl(self.root / "datasets/judge-framework-review/human_framework_reviews.jsonl", reviews)
        for paper_id in ("p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8", "p9"):
            _write_json(
                self.root / "datasets/judge-framework-all/evidence/structured" / (paper_id + ".json"),
                self._contract(paper_id, ["alpha_module", "beta_module"], ["alpha_module_to_beta_module"]),
            )
        _write_json(
            self.root / "datasets/judge-framework-all/evidence/structured/p7.json",
            self._contract("p7", [], []),
        )
        for paper_id in ("p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8", "p9"):
            pdf = self.root / "experiments/literature_review/papers/pdf" / (paper_id + ".pdf")
            pdf.parent.mkdir(parents=True, exist_ok=True)
            pdf.write_bytes(b"%PDF-1.4 fake")
        _png(self.root / "icon.png")
        for paper_id in ("p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8"):
            _png(self.root / "experiments/literature_review/papers/_stage" / paper_id / "cand_1.png")
        _png(self.root / "experiments/literature_review/papers/_stage/p9/cand_1.png")
        _png(self.root / "experiments/literature_review/papers/_reselect_good/p9/cand_1.png", "black")
        self._categories_csv()

    def prepare(self, target: int = 5, **kwargs):
        return bench.prepare_gold_instances(self.root / "gold", repo_root=self.root, target=target, **kwargs)


class LabelGroupFilterTests(FigureBenchFixture):
    """A release may be built from human-approved reference figures only."""

    def test_good_only_drops_every_poor_reference(self):
        report = self.prepare(target=10, label_groups=("good",))
        self.assertTrue(report["ok"], report)
        self.assertEqual(4, report["instance_count"])
        self.assertEqual({"good": 4}, report["label_counts"])
        self.assertEqual(2, report["excluded"]["label_group_excluded"])
        manifest = json.loads((self.root / "gold" / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(["good"], manifest["selection"]["label_groups"])
        self.assertEqual(4, len(manifest["instances"]))
        for entry in manifest["instances"]:
            paper = json.loads((self.root / "gold" / "instances" / entry["paper_id"] / "paper.json").read_text(encoding="utf-8"))
            self.assertEqual("good", paper["label_group"], entry["paper_id"])

    def test_the_default_still_allows_both_groups(self):
        report = self.prepare(target=10)
        self.assertEqual(4, report["label_counts"]["good"])
        self.assertEqual(2, report["label_counts"]["bad"])
        manifest = json.loads((self.root / "gold" / "manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(["good", "bad"], manifest["selection"]["label_groups"])
        self.assertNotIn("label_group_excluded", report["excluded"])

    def test_good_only_is_still_deterministic(self):
        self.prepare(target=10, label_groups=("good",))
        manifest = (self.root / "gold" / "manifest.json").read_bytes()
        second = self.prepare(target=10, label_groups=("good",))
        self.assertEqual(4, second["instance_count"])
        self.assertEqual(manifest, (self.root / "gold" / "manifest.json").read_bytes())


class DomainStrataTests(FigureBenchFixture):
    """Every frozen instance carries one arXiv-derived domain label."""

    def _manifest(self) -> dict:
        return json.loads((self.root / "gold/manifest.json").read_text(encoding="utf-8"))

    def test_freeze_records_a_domain_for_every_instance(self):
        report = self.prepare(target=10)
        self.assertTrue(report["ok"], report)
        manifest = self._manifest()
        self.assertTrue(manifest["domains"]["available"])
        self.assertEqual(domains.DOMAIN_SCHEMA_VERSION, manifest["domains"]["schema_version"])
        self.assertEqual([], manifest["domains"]["unlabeled"])
        self.assertEqual(len(manifest["instances"]), sum(manifest["domains"]["counts"].values()))
        self.assertEqual(report["domain_counts"], manifest["domains"]["counts"])
        for entry in manifest["instances"]:
            paper = json.loads((self.root / "gold" / "instances" / entry["paper_id"] / "paper.json").read_text(encoding="utf-8"))
            self.assertIn(paper["domain"], domains.DOMAIN_LABELS, entry["paper_id"])
            self.assertEqual(entry["domain"], paper["domain"], entry["paper_id"])
            self.assertFalse(paper["domain_unlabeled"], entry["paper_id"])

    def test_a_paper_missing_from_the_metadata_is_marked_unlabeled(self):
        self.prepare(target=10)
        entry = self._manifest()["instances"][0]
        kept = [
            {"arxiv_id": paper_id, "primary_category": "cs.LG", "categories": "cs.LG"}
            for paper_id in ("p1", "p2", "p3", "p4", "p5", "p6", "p7", "p8", "p9")
            if paper_id != entry["paper_id"]
        ]
        self._categories_csv(kept)
        self.prepare(target=10)
        manifest = self._manifest()
        self.assertTrue(manifest["domains"]["available"])
        self.assertIn(entry["paper_id"], manifest["domains"]["unlabeled"])
        paper = json.loads((self.root / "gold" / "instances" / entry["paper_id"] / "paper.json").read_text(encoding="utf-8"))
        self.assertIsNone(paper["domain"])
        self.assertTrue(paper["domain_unlabeled"])
        verification = bench.verify_gold_instances(self.root / "gold")
        self.assertTrue(verification["ok"], verification)

    def test_prepare_without_category_metadata_degrades_and_verify_warns(self):
        (self.root / "experiments/literature_review/papers/papers.csv").unlink()
        self.prepare(target=10)
        manifest = self._manifest()
        self.assertFalse(manifest["domains"]["available"])
        self.assertEqual(len(manifest["instances"]), len(manifest["domains"]["unlabeled"]))
        verification = bench.verify_gold_instances(self.root / "gold")
        self.assertTrue(verification["ok"], verification)
        self.assertIn("gold set carries no domain strata", verification["warnings"])

    def test_verify_flags_a_domain_label_that_drifted(self):
        self.prepare(target=10)
        entry = self._manifest()["instances"][0]
        path = self.root / "gold" / "instances" / entry["paper_id"] / "paper.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["domain"] = "vision" if entry["domain"] != "vision" else "agent"
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        verification = bench.verify_gold_instances(self.root / "gold")
        self.assertFalse(verification["ok"])
        self.assertTrue(any("domain label drifted" in error for error in verification["errors"]), verification)


class CohortTests(FigureBenchFixture):
    """The released freeze keeps held-out and Judge-overlapping papers apart."""

    def _manifest(self) -> dict:
        return json.loads((self.root / "gold/manifest.json").read_text(encoding="utf-8"))

    def test_the_default_freezes_only_the_held_out_cohort(self):
        report = self.prepare(target=5)
        self.assertEqual({"held_out": 5}, report["cohort_counts"])
        self.assertEqual(1, report["excluded"]["outside_held_out_splits"])
        for entry in self._manifest()["instances"]:
            paper = json.loads((self.root / "gold" / "instances" / entry["paper_id"] / "paper.json").read_text(encoding="utf-8"))
            self.assertEqual("held_out", entry["cohort"])
            self.assertEqual(entry["cohort"], paper["cohort"])

    def test_drawing_on_every_split_adds_the_overlap_cohort(self):
        report = self.prepare(target=10, include_splits=("train", "test", "validation"))
        self.assertEqual({"held_out": 6, "judge_train_overlap": 1}, report["cohort_counts"])
        self.assertNotIn("outside_held_out_splits", report["excluded"])
        manifest = self._manifest()
        self.assertEqual(["train", "test", "validation"], manifest["selection"]["include_splits"])
        trained = json.loads((self.root / "gold" / "instances" / "p5" / "paper.json").read_text(encoding="utf-8"))
        self.assertEqual("judge_train_overlap", trained["cohort"])
        self.assertIn("judge_train_overlap", manifest["cohorts"]["labels"])

    def test_verify_flags_a_cohort_that_drifted(self):
        self.prepare(target=5, include_splits=("train", "test", "validation"))
        path = self.root / "gold" / "instances" / "p5" / "paper.json"
        payload = json.loads(path.read_text(encoding="utf-8"))
        payload["cohort"] = "held_out"
        self.assertNotEqual("held_out", next(e for e in self._manifest()["instances"] if e["paper_id"] == "p5")["cohort"])
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        verification = bench.verify_gold_instances(self.root / "gold")
        self.assertFalse(verification["ok"])
        self.assertTrue(any("cohort drifted" in error for error in verification["errors"]), verification)


class GoldFreezeTests(FigureBenchFixture):
    def test_freeze_selects_held_out_good_first_and_reports_exclusions(self):
        report = self.prepare(target=5)
        self.assertTrue(report["ok"], report)
        self.assertEqual(4, report["label_counts"]["good"])
        self.assertEqual(1, report["label_counts"]["bad"])
        self.assertEqual(6, report["candidate_pool"])
        self.assertEqual(1, report["excluded"]["outside_held_out_splits"])
        self.assertEqual(1, report["excluded"]["review_flagged"])
        self.assertEqual(1, report["excluded"]["empty_contract"])

    def test_freeze_prefers_the_reselected_candidate(self):
        self.prepare(target=5)
        manifest = json.loads((self.root / "gold/manifest.json").read_text(encoding="utf-8"))
        entry = next(item for item in manifest["instances"] if item["paper_id"] == "p9")
        self.assertEqual("p9__reselection_good_1", entry["reference_figure_id"])
        self.assertEqual("reselection_good", json.loads((self.root / "gold/instances/p9/paper.json").read_text(encoding="utf-8"))["candidate_source"])

    def test_freeze_is_reproducible_and_order_independent(self):
        self.prepare(target=5)
        first = (self.root / "gold/manifest.json").read_bytes()
        second_report = self.prepare(target=5)
        self.assertEqual(first, (self.root / "gold/manifest.json").read_bytes())
        self.assertEqual(5, second_report["instance_count"])
        figures = (self.root / "datasets/judge-framework-v001/figures.jsonl").read_text(encoding="utf-8").splitlines()
        _write_jsonl(self.root / "datasets/judge-framework-v001/figures.jsonl", [json.loads(line) for line in reversed(figures)])
        self.assertEqual(first, (self.prepare(target=5), (self.root / "gold/manifest.json").read_bytes())[1])

    def test_freeze_keeps_hand_written_confirmation(self):
        self.prepare(target=5)
        confirmation = self.root / "gold/instances/p1/confirmation.json"
        payload = json.loads(confirmation.read_text(encoding="utf-8"))
        payload["reviewer"] = "reviewer-a"
        payload["checks"]["entities_complete"] = True
        _write_json(confirmation, payload)
        self.prepare(target=5)
        self.assertEqual("reviewer-a", json.loads(confirmation.read_text(encoding="utf-8"))["reviewer"])

    def test_verify_reports_a_changed_frozen_asset(self):
        self.prepare(target=5)
        self.assertTrue(bench.verify_gold_instances(self.root / "gold")["ok"])
        _png(self.root / "gold/instances/p1/reference.png", "red")
        result = bench.verify_gold_instances(self.root / "gold")
        self.assertFalse(result["ok"])
        self.assertTrue(any("frozen asset changed" in error for error in result["errors"]))

    def test_verify_reports_a_missing_instance_file(self):
        self.prepare(target=5)
        (self.root / "gold/instances/p2/reference.png").unlink()
        result = bench.verify_gold_instances(self.root / "gold")
        self.assertFalse(result["ok"])
        self.assertTrue(any("missing instance file" in error for error in result["errors"]))

    def test_verify_reports_a_missing_gold_root(self):
        result = bench.verify_gold_instances(self.root / "absent")
        self.assertFalse(result["ok"])
        self.assertTrue(result["errors"])

    def test_refuses_to_overwrite_a_different_asset_without_force(self):
        self.prepare(target=5)
        source = self.root / "experiments/literature_review/papers/_stage/p1/cand_1.png"
        _png(source, "red")
        with self.assertRaises(ValueError):
            self.prepare(target=5)
        self.prepare(target=5, force=True)


class MatchingTests(unittest.TestCase):
    def test_strict_needs_the_same_identifier_but_soft_accepts_extra_words(self):
        gold = ["attention_mechanism"]
        predicted = ["attention mechanism module"]
        self.assertEqual(0.0, matching.match_entities(gold, predicted, mode="strict")["f1"])
        self.assertEqual(1.0, matching.match_entities(gold, predicted, mode="soft")["f1"])

    def test_normalisation_collapses_identifiers_and_human_labels(self):
        self.assertEqual(
            matching.normalize_key("exploration_probability_alpha"),
            matching.normalize_key("Exploration probability alpha"),
        )
        report = matching.match_entities(["exploration_probability_alpha"], ["Exploration probability alpha"], mode="strict")
        self.assertEqual(1.0, report["f1"])

    def test_precision_recall_are_symmetric_for_a_partial_overlap(self):
        gold = ["alpha_module", "beta_module"]
        predicted = ["alpha module", "gamma module"]
        report = matching.match_entities(gold, predicted, mode="soft")
        self.assertEqual(1, report["matched"])
        self.assertEqual(0.5, report["precision"])
        self.assertEqual(0.5, report["recall"])
        self.assertEqual(0.5, report["f1"])

    def test_direction_is_never_relaxed(self):
        gold = [{"source": "alpha", "target": "beta", "relation_type": "feeds"}]
        reversed_edge = [{"source": "beta", "target": "alpha", "relation_type": "feeds"}]
        self.assertEqual(0, matching.match_relations(gold, reversed_edge, mode="soft")["matched"])

    def test_untyped_gold_relation_matches_on_endpoints(self):
        gold = ["alpha_module_to_beta_module"]
        parsed = bench._resolve_relation_entry("alpha_module_to_beta_module", {}, ["alpha_module", "beta_module"])
        self.assertEqual({"source": "alpha_module", "target": "beta_module", "relation_type": ""}, parsed)
        predicted = [{"source": "alpha_module", "target": "beta_module", "relation_type": "data_flow"}]
        self.assertEqual(1, matching.match_relations([parsed], predicted, mode="strict")["matched"])

    def test_relation_type_aliases_collapse_wording_but_not_direction(self):
        self.assertEqual("feed", matching.normalize_relation_type("provides input to"))
        self.assertEqual("feed", matching.normalize_relation_type("feeds into"))
        gold = [{"source": "alpha", "target": "beta", "relation_type": "provides input to"}]
        predicted = [{"source": "alpha", "target": "beta", "relation_type": "feeds"}]
        self.assertEqual(0.0, matching.match_relations(gold, predicted, mode="strict")["f1"])
        self.assertEqual(1.0, matching.match_relations(gold, predicted, mode="soft")["f1"])

    def test_greek_letters_normalise_to_their_names(self):
        self.assertEqual(matching.normalize_key("exploration_probability_alpha"), matching.normalize_key("exploration probability \u03b1"))

    def test_rendered_recall_reads_the_delivered_text(self):
        report = matching.rendered_entity_recall(["alpha_module", "beta_module"], "The Alpha module drives the beta module")
        self.assertEqual(2, report["matched"])
        self.assertEqual(1.0, report["recall"])


class ModuleCoverageTests(unittest.TestCase):
    """The L1 ruler reads coverage over named modules and splits notation out."""

    def test_notation_is_separated_from_named_modules(self):
        modules, notation = matching.split_notation(
            ["agent_policy", "o_{-i}", "Mask(X)", "llm", "g_{att}(o_t)", "Decision Tree", "X"]
        )
        self.assertEqual(["agent_policy", "llm", "Decision Tree"], modules)
        self.assertEqual(["o_{-i}", "Mask(X)", "g_{att}(o_t)", "X"], notation)

    def test_short_acronyms_stay_modules_rather_than_notation(self):
        for label in ("llm", "vlm", "gru", "prd"):
            self.assertFalse(matching.is_notation(label), label)

    def test_notation_labels_do_not_take_a_module_slot(self):
        report = matching.module_coverage(["alpha_module"], ["o_{-i}", "alpha module"], mode="strict")
        self.assertEqual(1, report["matched"])
        self.assertEqual(1.0, report["coverage"])
        self.assertEqual(1, report["predicted_notation_count"])

    def test_soft_matching_treats_single_character_tokens_as_indices(self):
        report = matching.match_entities(["exploration_probability_alpha"], ["\\alpha_i"], mode="soft")
        self.assertEqual(1, report["matched"])

    def test_a_label_without_a_longer_token_still_binds_nothing(self):
        report = matching.match_entities(["alpha_module"], ["x"], mode="soft")
        self.assertEqual(0, report["matched"])

    def test_notation_can_represent_a_module_named_after_a_symbol(self):
        report = matching.module_coverage(["exploration_probability_alpha"], ["\\alpha_i"], mode="soft")
        self.assertEqual(1, report["matched"])
        self.assertEqual(1.0, report["coverage"])
        self.assertEqual(1, report["notation_modules_matched"])

    def test_unbound_notation_is_reported_but_never_counted_as_a_drawn_module(self):
        report = matching.unbound_nodes(
            ["exploration_probability_alpha"], ["\\alpha_i", "\\hat{a}_t"], mode="soft"
        )
        self.assertEqual(0, report["unbound"])
        self.assertEqual(1, report["predicted_module_count"])
        self.assertEqual(1, report["notation_unbound_excluded"])

    def test_coverage_counts_represented_modules_and_lists_the_missing(self):
        report = matching.module_coverage(
            ["alpha_module", "beta_module", "gamma_module"],
            ["alpha module", "beta module"],
            mode="strict",
        )
        self.assertEqual(2, report["matched"])
        self.assertEqual(3, report["gold_module_count"])
        self.assertEqual(0.6667, report["coverage"])
        self.assertEqual(["gamma_module"], report["missing"])

    def test_unbound_rate_counts_drawn_modules_the_contract_never_asked_for(self):
        report = matching.unbound_nodes(["alpha_module"], ["alpha module", "extra_panel"], mode="strict")
        self.assertEqual(1, report["unbound"])
        self.assertEqual(2, report["predicted_module_count"])
        self.assertEqual(0.5, report["rate"])
        self.assertEqual(["extra_panel"], report["labels"])


class GoldRelationReadingTests(unittest.TestCase):
    def test_reads_objects_ids_pairs_and_free_text(self):
        contract = {
            "relations": [{"source": "alpha_module", "target": "beta_module", "relation_type": "feeds"}],
            "entities": [{"entity_id": "alpha_module"}, {"entity_id": "beta_module"}, {"entity_id": "gamma_module"}],
        }
        content = {
            "required_entities": ["alpha_module", "beta_module", "gamma_module"],
            "required_relations": [
                {"source": "alpha_module", "target": "beta_module", "relation_type": "feeds"},
                "alpha_module_to_beta_module",
                "alpha_module->beta_module",
                ["alpha_module", "beta_module"],
                "alpha_module contains beta_module",
                "alpha_module-feeds-beta_module",
                "gamma_module",
            ],
        }
        reading = bench._gold_relations(content, contract)
        self.assertEqual(6, len(reading["relations"]))
        self.assertEqual(1, reading["unresolved"])
        self.assertEqual(7, reading["declared"])
        types = {item["relation_type"] for item in reading["relations"]}
        self.assertIn("contains", types)
        self.assertIn("feeds", types)


class NativeEditabilityTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.picture = self.root / "icon.png"
        _png(self.picture)

    def tearDown(self):
        self._tmp.cleanup()

    def test_counts_connectors_text_and_raster_pictures(self):
        deck = self.root / "deck.pptx"
        _deck_with_connector(deck, self.picture)
        report = measure_native_editability(deck)
        self.assertTrue(report["ok"])
        self.assertEqual(1, report["native_connector_count"])
        self.assertEqual(0, report["line_shape_count"])
        self.assertEqual(1, report["arrow_capable_shape_count"])
        self.assertEqual(4, report["leaf_shape_count"])
        self.assertEqual(0, report["group_count"])
        self.assertEqual(2, report["autoshape_count"])
        # Both the text box and the rounded rectangle carry an editable text frame,
        # even though only one of them holds characters.
        self.assertEqual(2, report["text_frame_count"])
        self.assertEqual(len("Alpha module"), report["editable_text_char_count"])
        self.assertEqual(1, report["picture_count"])
        self.assertEqual(1, report["raster_picture_count"])
        self.assertEqual(0, report["vector_picture_count"])
        self.assertTrue(report["has_connectors"])
        self.assertLess(report["raster_area_ratio"], 1.0)
        self.assertIsNone(report["native_connector_ratio"])
        self.assertIsNone(report["connector_realization_ratio"])
        self.assertTrue(report["ooxml_valid"])

    def test_realization_ratio_uses_the_gold_relation_count(self):
        deck = self.root / "deck.pptx"
        _deck_with_connector(deck, self.picture)
        report = measure_native_editability(deck, expected_connector_count=4)
        self.assertEqual(0.25, report["connector_realization_ratio"])

    def test_text_is_returned_only_when_requested(self):
        deck = self.root / "deck.pptx"
        _deck_with_connector(deck, self.picture, text="Editable label")
        self.assertNotIn("editable_text", measure_native_editability(deck))
        self.assertIn("Editable label", measure_native_editability(deck, include_text=True)["editable_text"])

    def test_missing_deck_fails_explicitly(self):
        report = measure_native_editability(self.root / "absent.pptx")
        self.assertFalse(report["ok"])
        self.assertFalse(report["ooxml_valid"])
        self.assertEqual("missing pptx", report["error"])

    def test_corrupt_deck_fails_explicitly(self):
        broken = self.root / "broken.pptx"
        broken.write_bytes(b"not a deck")
        report = measure_native_editability(broken)
        self.assertFalse(report["ok"])
        self.assertFalse(report["ooxml_valid"])
        self.assertIn("error", report)


class InstanceScoringTests(FigureBenchFixture):
    def _gold_instance(self, paper_id: str = "p1"):
        self.prepare(target=5)
        return self.root / "gold/instances" / paper_id

    def _run_dir(self, name: str, nodes, arrows, banned_label: str | None = None) -> Path:
        run = self.root / "systems" / name
        _write_json(
            run / "figure_program.json",
            {
                "canvas": {"width_in": 13.333, "height_in": 7.5},
                "semantic_nodes": [{"id": node, "label": node.replace("_", " ")} for node in nodes],
                "arrows": [{"id": "rel", "source": source, "target": target, "type": "feeds"} for source, target in arrows],
                "text_program": {"items": [{"id": "txt", "text": banned_label or "alpha module"}]},
            },
        )
        _deck_with_connector(run / "editable_composition.pptx", self.root / "icon.png", text="alpha module")
        return run

    def test_matching_program_scores_perfect_content_fidelity(self):
        instance = self._gold_instance()
        run = self._run_dir("system-ok", ["alpha_module", "beta_module"], [("alpha_module", "beta_module")])
        metrics = bench.evaluate_instance(instance, run)
        self.assertEqual(1.0, metrics["L1"]["entity_strict"]["f1"])
        self.assertEqual(1.0, metrics["L1"]["relation_soft"]["f1"])
        self.assertFalse(metrics["hard_gates"]["failed"])
        self.assertGreaterEqual(metrics["L2"]["native_connector_count"], 1)
        self.assertEqual(2, metrics["gold"]["entity_count"])

    def test_invented_entity_trips_the_hard_gate(self):
        instance = self._gold_instance()
        run = self._run_dir("system-inventing", ["alpha_module", "beta_module"], [], banned_label="centralized execution module")
        metrics = bench.evaluate_instance(instance, run)
        self.assertTrue(metrics["hard_gates"]["failed"])
        self.assertIn("must_not_invent_hit", metrics["hard_gates"]["failed_gates"])
        self.assertEqual(["centralized_execution_module"], metrics["hard_gates"]["must_not_invent_hits"])

    def test_missing_deck_trips_the_hard_gate(self):
        instance = self._gold_instance()
        run = self.root / "systems/system-empty"
        run.mkdir(parents=True, exist_ok=True)
        metrics = bench.evaluate_instance(instance, run)
        self.assertEqual(["missing_editable_composition"], metrics["hard_gates"]["failed_gates"])

    def test_unreadable_deck_trips_the_hard_gate(self):
        instance = self._gold_instance()
        run = self._run_dir("system-broken", ["alpha_module"], [])
        (run / "editable_composition.pptx").write_bytes(b"not a deck")
        metrics = bench.evaluate_instance(instance, run)
        self.assertIn("unreadable_powerpoint", metrics["hard_gates"]["failed_gates"])

    def test_layers_can_be_limited(self):
        instance = self._gold_instance()
        run = self._run_dir("system-l2", ["alpha_module"], [])
        metrics = bench.evaluate_instance(instance, run, layers=("L2",))
        self.assertIn("L2", metrics)
        self.assertNotIn("L1", metrics)
        self.assertNotIn("L3", metrics)


    def test_l1_reports_module_coverage_and_unbound_nodes(self):
        instance = self._gold_instance()
        run = self._run_dir("system-notation", ["alpha_module", "beta_module", "Mask(X)"], [("alpha_module", "beta_module")])
        block = bench.evaluate_instance(instance, run)["L1"]["core_modules"]
        self.assertEqual(2, block["gold_module_count"])
        self.assertEqual(1.0, block["coverage"]["strict"]["coverage"])
        self.assertEqual(0, block["unbound"]["strict"]["unbound"])
        self.assertEqual(1, block["predicted_notation_count"])
        self.assertEqual(0.3333, block["notation_ratio"])


class RunAndScoreTests(FigureBenchFixture):
    def _systems(self) -> Path:
        self.prepare(target=5)
        instance_ids = sorted(child.name for child in (self.root / "gold/instances").iterdir())
        for name in ("system-ok", "system-empty"):
            for paper_id in instance_ids:
                run = self.root / "systems" / name / paper_id
                run.mkdir(parents=True, exist_ok=True)
                if name == "system-ok":
                    _write_json(
                        run / "figure_program.json",
                        {
                            "canvas": {"width_in": 13.333, "height_in": 7.5},
                            "semantic_nodes": [{"id": "alpha_module", "label": "alpha module"}],
                            "arrows": [],
                            "text_program": {"items": []},
                        },
                    )
                    _deck_with_connector(run / "editable_composition.pptx", self.root / "icon.png")
        return self.root / "systems"

    def test_run_writes_one_metrics_file_per_instance_and_system(self):
        systems = self._systems()
        report = bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems,
            out_dir=self.root / "runs",
        )
        self.assertEqual(2, len(report["systems"]))
        ok = next(item for item in report["systems"] if item["system_id"] == "system-ok")
        empty = next(item for item in report["systems"] if item["system_id"] == "system-empty")
        self.assertEqual(5, ok["scored_instances"])
        self.assertEqual(0, ok["hard_gate_failures"])
        self.assertEqual(5, empty["hard_gate_failures"])
        self.assertTrue((self.root / "runs/system-ok/p1/metrics.json").is_file())

    def test_run_accepts_a_single_system_directory(self):
        systems = self._systems()
        report = bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems / "system-ok",
            out_dir=self.root / "runs-single",
            system_id="chosen",
        )
        self.assertEqual(["chosen"], [item["system_id"] for item in report["systems"]])

    def test_score_reports_a_domain_breakdown(self):
        systems = self._systems()
        bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems,
            out_dir=self.root / "runs",
        )
        report = bench.score_benchmark(runs_dir=self.root / "runs", out_dir=self.root / "score")
        summary = report["systems"]["system-ok"]
        self.assertTrue(summary["by_domain"], summary)
        for domain, layers in summary["by_domain"].items():
            self.assertIn(domain, domains.DOMAIN_LABELS, domain)
            self.assertIn("L1", layers)
            self.assertGreater(layers["L1"]["entity_strict_f1"]["count"], 0)
        strata = {
            json.loads((self.root / "runs/system-ok" / paper_id / "metrics.json").read_text(encoding="utf-8"))["gold"]["domain"]
            for paper_id in sorted(child.name for child in (self.root / "gold/instances").iterdir())
        }
        self.assertEqual(sorted(strata), sorted(summary["by_domain"]))
        table = (self.root / "score/score_table.md").read_text(encoding="utf-8")
        self.assertIn("domain:", table)

    def test_score_reports_a_cohort_breakdown(self):
        systems = self._systems()
        bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems,
            out_dir=self.root / "runs",
        )
        report = bench.score_benchmark(runs_dir=self.root / "runs", out_dir=self.root / "score")
        cohorts = report["systems"]["system-ok"]["by_cohort"]
        self.assertEqual(["held_out"], sorted(cohorts))
        self.assertEqual(5, cohorts["held_out"]["L1"]["entity_strict_f1"]["count"])
        table = (self.root / "score/score_table.md").read_text(encoding="utf-8")
        self.assertIn("cohort: held_out", table)

    def test_score_table_reports_module_coverage_and_unbound_rate(self):
        systems = self._systems()
        bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems,
            out_dir=self.root / "runs",
        )
        bench.score_benchmark(runs_dir=self.root / "runs", out_dir=self.root / "score")
        table = (self.root / "score/score_table.md").read_text(encoding="utf-8")
        self.assertIn("module_coverage_soft", table)
        self.assertIn("unbound_node_rate_soft", table)
        self.assertIn("module_coverage_strict", table)

    def test_score_is_byte_identical_on_a_rerun(self):
        systems = self._systems()
        bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems,
            out_dir=self.root / "runs",
        )
        first = bench.score_benchmark(runs_dir=self.root / "runs", out_dir=self.root / "score")
        table = (self.root / "score/score_table.md").read_bytes()
        payload = (self.root / "score/score_report.json").read_bytes()
        bench.score_benchmark(runs_dir=self.root / "runs", out_dir=self.root / "score")
        self.assertEqual(table, (self.root / "score/score_table.md").read_bytes())
        self.assertEqual(payload, (self.root / "score/score_report.json").read_bytes())
        self.assertEqual(hashlib.sha256(table).hexdigest(), hashlib.sha256((self.root / "score/score_table.md").read_bytes()).hexdigest())

    def test_score_reports_gate_failures_and_per_quality_breakdown(self):
        systems = self._systems()
        bench.run_benchmark_instances(
            instances_dir=self.root / "gold/instances",
            systems_dir=systems,
            out_dir=self.root / "runs",
        )
        report = bench.score_benchmark(runs_dir=self.root / "runs", out_dir=self.root / "score")
        empty = report["systems"]["system-empty"]
        self.assertEqual(0.0, empty["hard_gate_pass_rate"])
        self.assertEqual(5, empty["hard_gate_failures"]["missing_editable_composition"])
        self.assertEqual({"good", "bad"}, set(empty["by_reference_quality"]))
        self.assertIn("L1", empty["layers"])
        table = (self.root / "score/score_table.md").read_text(encoding="utf-8")
        self.assertIn("reference quality: good", table)

    def test_score_rejects_an_empty_run_tree(self):
        (self.root / "empty-runs").mkdir(parents=True, exist_ok=True)
        with self.assertRaises(ValueError):
            bench.score_benchmark(runs_dir=self.root / "empty-runs", out_dir=self.root / "score")

    def test_run_rejects_a_missing_systems_directory(self):
        self.prepare(target=5)
        with self.assertRaises(ValueError):
            bench.run_benchmark_instances(
                instances_dir=self.root / "gold/instances",
                systems_dir=self.root / "absent",
                out_dir=self.root / "runs",
            )


class CommandLineTests(unittest.TestCase):
    def test_benchmark_actions_are_registered(self):
        from rfs.cli import build_parser

        parser = build_parser()
        for action in ("prepare", "run", "score"):
            namespace = parser.parse_args(["benchmark", action])
            self.assertEqual(action, namespace.benchmark_action)
        namespace = parser.parse_args(["benchmark", "run", "--instances", "a", "--systems", "b", "--out", "c", "--layers", "L2"])
        self.assertEqual("L2", namespace.layers)
        self.assertEqual("L1,L2,L3", parser.parse_args(["benchmark", "score"]).layers)

    def test_unknown_layer_is_rejected(self):
        from rfs.cli import _benchmark_layers, build_parser

        namespace = build_parser().parse_args(["benchmark", "score", "--layers", "L1,L9"])
        with self.assertRaises(ValueError):
            _benchmark_layers(namespace)

    def test_benchmark_run_requires_its_inputs(self):
        from rfs.cli import _dispatch, build_parser

        namespace = build_parser().parse_args(["benchmark", "run"])
        with self.assertRaises(ValueError):
            _dispatch(namespace)

    def test_score_own_requires_a_manifest_and_does_not_score_an_empty_delivery(self):
        from rfs.cli import _dispatch, build_parser

        parser = build_parser()
        with self.assertRaises(ValueError):
            _dispatch(parser.parse_args(["benchmark", "score-own"]))
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = root / "manifest.jsonl"
            manifest.write_text('{"paper_id": "doi/a"}\n', encoding="utf-8")
            report = _dispatch(parser.parse_args(["benchmark", "score-own", "--manifest", str(manifest), "--runs", str(root / "runs"), "--paper-quiz", "off", "--nine-judge", "off"]))
        self.assertEqual(1, report["delivery_missing"])
        self.assertEqual(0, report["scored"])
        self.assertFalse((root / "runs" / "doi_a" / "metrics.json").exists())


if __name__ == "__main__":
    unittest.main()


class FakeEmbedder:
    """Deterministic stand-in for a captured text embedder, so no test hits the network."""

    def __init__(self, table):
        self.table = table
        self.seen = []

    def encode(self, texts):
        self.seen.extend(texts)
        return [list(self.table[text]) for text in texts]


class EmbeddingLevelTests(unittest.TestCase):
    """The optional L1 embedding level: opt-in, published threshold, honest absence."""

    def test_disabled_without_the_environment_variable(self):
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop(matching.EMBEDDING_ENV_VAR, None)
            embedder, spec, reason = bench.resolve_embedder()
        self.assertIsNone(embedder)
        self.assertEqual("", spec)
        self.assertEqual("disabled", reason)

    def test_the_environment_variable_opts_the_level_in(self):
        with mock.patch.dict(os.environ, {matching.EMBEDDING_ENV_VAR: "api:unit-test-embedding"}, clear=False):
            embedder, spec, reason = bench.resolve_embedder()
        self.assertIsNotNone(embedder)
        self.assertEqual("api:unit-test-embedding", spec)
        self.assertEqual("", reason)
        self.assertEqual("unit-test-embedding", embedder.fingerprint()["model"])

    def test_the_embedder_is_shared_per_model_so_its_cache_survives_the_run(self):
        first = bench.resolve_embedder("api:unit-test-embedding")[0]
        second = bench.resolve_embedder("api:unit-test-embedding")[0]
        self.assertIs(first, second)

    def test_an_unsupported_spec_is_reported_rather_than_guessed(self):
        embedder, spec, reason = bench.resolve_embedder("local:hashing-ngram")
        self.assertIsNone(embedder)
        self.assertEqual("local:hashing-ngram", spec)
        self.assertEqual("unsupported_spec", reason)

    def test_provenance_records_absence_with_a_reason(self):
        block = bench._embedder_provenance(None, "", "disabled")
        self.assertFalse(block["used"])
        self.assertEqual("disabled", block["reason"])
        self.assertEqual(matching.EMBEDDING_ENV_VAR, block["env_var"])

    def test_embedding_mode_matches_a_paraphrase_the_token_sets_miss(self):
        embedder = FakeEmbedder({"sparse_reward_signal": [1.0, 0.0], "the reward sparsity term": [0.99, 0.01]})
        strict = matching.match_entities(["sparse_reward_signal"], ["the reward sparsity term"], mode="strict")
        embedding = matching.match_entities(
            ["sparse_reward_signal"], ["the reward sparsity term"], mode="embedding", embedder=embedder
        )
        self.assertEqual(0.0, strict["f1"])
        self.assertEqual(1.0, embedding["f1"])
        self.assertEqual("embedding", embedding["mode"])

    def test_embedding_mode_rejects_a_pair_below_the_published_threshold(self):
        embedder = FakeEmbedder({"alpha_module": [1.0, 0.0], "beta_module": [0.0, 1.0]})
        report = matching.match_entities(["alpha_module"], ["beta_module"], mode="embedding", embedder=embedder)
        self.assertEqual(0.0, report["f1"])
        self.assertGreater(matching.EMBEDDING_MATCH_THRESHOLD, 0.0)

    def test_embedding_mode_without_an_embedder_is_an_error(self):
        with self.assertRaises(ValueError):
            matching.match_entities(["alpha"], ["alpha"], mode="embedding")

    def test_cosine_is_zero_for_degenerate_or_mismatched_vectors(self):
        self.assertEqual(0.0, matching.cosine_similarity([0.0, 0.0], [1.0, 0.0]))
        self.assertEqual(0.0, matching.cosine_similarity([], []))
        self.assertEqual(0.0, matching.cosine_similarity([1.0, 0.0], [1.0, 0.0, 0.0]))
        self.assertEqual(1.0, matching.cosine_similarity([1.0, 0.0], [1.0, 0.0]))

    def test_the_offline_levels_ignore_an_embedder(self):
        embedder = FakeEmbedder({"alpha module": [1.0, 0.0]})
        report = matching.match_entities(["alpha_module"], ["alpha module"], mode="soft", embedder=embedder)
        self.assertEqual(1.0, report["f1"])
        self.assertEqual([], embedder.seen)


class DeckDerivedPredictionsTests(unittest.TestCase):
    """A system that ships no semantic program is still scored on what it drew."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        _png(self.root / "icon.png")

    def tearDown(self):
        self._tmp.cleanup()

    def test_text_frames_become_entity_predictions(self):
        deck = self.root / "editable_composition.pptx"
        _deck_with_connector(deck, self.root / "icon.png", text="alpha module")
        block = bench._deck_predictions(deck, ["alpha_module"])
        self.assertTrue(block["available"])
        self.assertEqual("deck_text_frames", block["source"])
        self.assertEqual(1, block["predicted_entity_count"])
        self.assertEqual(1.0, block["entity_soft"]["f1"])

    def test_relations_are_reported_unavailable_rather_than_guessed(self):
        deck = self.root / "editable_composition.pptx"
        _deck_with_connector(deck, self.root / "icon.png")
        block = bench._deck_predictions(deck, ["alpha_module"])
        self.assertFalse(block["relations"]["available"])
        self.assertIn("stCxn", block["relations"]["reason"])

    def test_a_missing_deck_is_reported_not_scored(self):
        block = bench._deck_predictions(self.root / "absent.pptx", ["alpha_module"])
        self.assertFalse(block["available"])
        self.assertIn("missing", block["reason"])

    def test_a_failing_embedder_is_recorded_instead_of_raised(self):
        deck = self.root / "editable_composition.pptx"
        _deck_with_connector(deck, self.root / "icon.png")

        class Broken:
            def encode(self, texts):
                raise RuntimeError("no credential")

        block = bench._deck_predictions(deck, ["alpha_module"], Broken())
        self.assertFalse(block["entity_embedding"]["available"])
        self.assertIn("no credential", block["entity_embedding"]["reason"])

    def test_shape_labels_skip_groups_and_empty_frames(self):
        deck = self.root / "labels.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2.0), Inches(0.5)).text_frame.text = "First label"
        slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(2.0), Inches(0.5)).text_frame.text = "   "
        slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(2.0), Inches(0.5)).text_frame.text = "Second label\nsecond line"
        deck.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(deck))
        self.assertEqual(["First label", "Second label"], bench.extract_shape_labels(deck))


class DeckOnlyInstanceScoringTests(FigureBenchFixture):
    """Scoring a deck-only system: L1 falls back to the deck, the gates still apply."""

    def _gold_instance(self, paper_id: str = "p1"):
        self.prepare(target=5)
        return self.root / "gold/instances" / paper_id

    def _deck_only_run(self, name: str, text: str = "alpha module") -> Path:
        run = self.root / "systems" / name
        _deck_with_connector(run / "editable_composition.pptx", self.root / "icon.png", text=text)
        return run

    def test_a_run_without_a_program_still_reports_l1_and_l2(self):
        instance = self._gold_instance()
        run = self._deck_only_run("deck-only")
        metrics = bench.evaluate_instance(instance, run)
        self.assertTrue(metrics["L1"]["deck"]["available"])
        self.assertEqual(0, metrics["L1"]["predicted_entity_count"])
        self.assertGreaterEqual(metrics["L2"]["native_connector_count"], 1)
        self.assertFalse(metrics["hard_gates"]["gates"]["missing_editable_composition"])

    def test_embedding_provenance_is_recorded_when_the_level_is_off(self):
        instance = self._gold_instance()
        run = self._deck_only_run("deck-only-2")
        with mock.patch.dict(os.environ, {}, clear=False):
            os.environ.pop(matching.EMBEDDING_ENV_VAR, None)
            metrics = bench.evaluate_instance(instance, run)
        self.assertFalse(metrics["L1"]["embedder"]["used"])
        self.assertEqual("disabled", metrics["L1"]["embedder"]["reason"])
        self.assertNotIn("entity_embedding", metrics["L1"])
        self.assertNotIn("entity_embedding", metrics["L1"]["deck"])

    def test_a_failing_embedding_call_is_recorded_and_the_rest_still_scores(self):
        instance = self._gold_instance()
        run = self._deck_only_run("deck-only-3")
        broken = mock.Mock()
        broken.encode.side_effect = RuntimeError("endpoint down")
        broken.fingerprint.return_value = {"model": "broken"}
        broken.usage.return_value = {"calls": 0}
        with mock.patch.object(bench, "resolve_embedder", return_value=(broken, "api:broken", "")):
            metrics = bench.evaluate_instance(instance, run)
        self.assertIn("embedding_failed", metrics["L1"]["embedder"]["reason"])
        self.assertIn("endpoint down", metrics["L1"]["embedder"]["reason"])
        self.assertFalse(metrics["L1"]["embedder"]["used"])
        self.assertFalse(metrics["L1"]["deck"]["entity_embedding"]["available"])
        self.assertTrue(metrics["L1"]["deck"]["available"])
