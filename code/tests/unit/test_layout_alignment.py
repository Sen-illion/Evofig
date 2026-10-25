"""Tests for reference-free PPTX alignment and spacing."""

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches

from rfs.evaluation.layout_alignment import measure_layout_alignment


def _row(path: Path, *, overflow: bool = False, tight: bool = False, uneven: bool = False, stagger: bool = False) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    width = Inches(1.2)
    height = Inches(0.8)
    gaps = [0.2, 0.9, 1.6] if uneven else [0.4, 0.4, 0.4]
    cursor = Inches(0.05) if tight else Inches(0.8)
    if overflow:
        cursor = Emu(-200000)
    for index in range(4):
        top = Inches(1.2 + index * 0.7) if stagger else Inches(2.0)
        slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cursor, top, width, height)
        if index < 3:
            cursor = Emu(int(cursor) + int(width) + int(Inches(gaps[index])))
    slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.8), Inches(3.2), Inches(4.0), Inches(3.2))
    path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(path)


class LayoutAlignmentTests(unittest.TestCase):
    def test_even_row_is_aligned_and_stays_on_the_canvas(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "even.pptx"
            _row(path)
            report = measure_layout_alignment(path)
        self.assertTrue(report["ok"])
        self.assertEqual(4, report["box_count"])
        self.assertGreater(report["aligned_edge_ratio"], 0.4)
        self.assertEqual(0.0, report["gap_cv"])
        self.assertEqual(0, report["overflow_count"])
        self.assertEqual(0, report["tight_margin_count"])

    def test_uneven_overflowing_row_scores_worse_geometry(self):
        with tempfile.TemporaryDirectory() as tmp:
            even = Path(tmp) / "even.pptx"
            messy = Path(tmp) / "messy.pptx"
            _row(even)
            _row(messy, overflow=True, uneven=True, stagger=True)
            even_report = measure_layout_alignment(even)
            messy_report = measure_layout_alignment(messy)
        self.assertLess(messy_report["aligned_edge_ratio"], even_report["aligned_edge_ratio"])
        self.assertGreater(messy_report["gap_cv"], 0)
        self.assertGreater(messy_report["overflow_count"], 0)

    def test_a_box_against_the_edge_is_tight_without_counting_as_overflow(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "tight.pptx"
            _row(path, tight=True)
            report = measure_layout_alignment(path)
        self.assertEqual(0, report["overflow_count"])
        self.assertGreater(report["tight_margin_count"], 0)

    def test_a_missing_deck_does_not_look_like_a_clean_figure(self):
        report = measure_layout_alignment(Path("missing-editable-composition.pptx"))
        self.assertFalse(report["ok"])
        self.assertIsNone(report["aligned_edge_ratio"])
        self.assertIsNone(report["overflow_count"])
        self.assertIn("missing", report["reason"])


if __name__ == "__main__":
    unittest.main()
