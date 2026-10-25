"""Construction of the public benchmark manifest, drafts, and score columns."""

import json
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from rfs.evaluation.benchmark_contracts import draft_contracts, index_structured_contracts
from rfs.evaluation.benchmark_manifest import (
    build_manifest_rows,
    export_public_benchmark_manifest,
    framework_figure_label,
    overlap_with_ids,
)
from rfs.evaluation.editable_figure_bench import evaluate_instance
from rfs.evaluation.external_benchmark_scores import figurebench_record, sciflow_record
from rfs.evaluation.livefigure_nine import parse_nine_scores, score_nine
from rfs.evaluation.native_editability import editability_verdict, measure_native_editability
from rfs.evaluation.paper_figure_clip import score_paper_figure_clip
from rfs.evaluation.own_set_score import method_section_from_text, score_own_delivery, score_own_manifest
from rfs.evaluation.paper_quiz import density_augmented_score, parse_paper_quiz, score_answer_bank, score_paper_quiz


def _jsonl(path: Path, rows) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(json.dumps(row) + "\n" for row in rows), encoding="utf-8")


class ManifestTests(unittest.TestCase):
    def test_framework_label_uses_only_a_human_decision(self):
        reviews = {"2301.02083": {"decision": "selected"}, "2301.00001": {"decision": "none"}}
        self.assertEqual("yes", framework_figure_label(["2301.02083"], reviews))
        self.assertEqual("no", framework_figure_label(["2301.00001"], reviews))
        self.assertEqual("unknown", framework_figure_label(["9999.00001"], reviews))

    def test_manifest_rows_keep_hidden_test_and_blank_missing_files(self):
        ledger = [
            {
                "paper_id": "10.1/hidden",
                "split": "hidden_test",
                "primary_domain": "cv",
                "figure_asset_tier": "A_source_verified",
                "figure_asset_sha256": ["abc"],
                "identity_keys": {"arxiv_id": ["2301.02083"]},
            },
            {"paper_id": "10.1/train", "split": "train", "identity_keys": {"arxiv_id": ["1.1"]}},
        ]
        rows = build_manifest_rows(
            ledger,
            snapshot=Path("missing-snapshot"),
            reviews_by_arxiv={"2301.02083": {"decision": "selected"}},
            framework_arxiv_ids={"2301.02083"},
            assets={},
            exclude_paper_ids={"10.1/train"},
        )
        self.assertEqual(["10.1/hidden"], [row["paper_id"] for row in rows])
        self.assertEqual("", rows[0]["pdf_path"])
        self.assertEqual("", rows[0]["official_figure_path"])
        self.assertEqual("yes", rows[0]["framework_figure"])
        self.assertTrue(rows[0]["in_framework_corpus"])
        self.assertIn("no longer a sealed Judge exam", rows[0]["public_benchmark_note"])

    def test_overlap_is_not_downloaded_when_the_id_file_is_absent(self):
        report = overlap_with_ids([{"paper_id": "a", "arxiv_id": "1"}], None, label="SciFlow-Bench")
        self.assertEqual("not_downloaded", report["status"])
        self.assertIsNone(report["overlap_count"])

    def test_export_appends_an_access_event_and_writes_overlap(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            _jsonl(
                root / "ledger.jsonl",
                [
                    {
                        "paper_id": "p",
                        "split": "hidden_test",
                        "primary_domain": "rag",
                        "figure_asset_tier": "B_official_standalone",
                        "identity_keys": {"arxiv_id": ["2305.1"]},
                    }
                ],
            )
            _jsonl(root / "reviews.jsonl", [{"arxiv_id": "2305.1", "decision": "none"}])
            _jsonl(root / "framework.jsonl", [])
            (root / "snapshot").mkdir()
            access = root / "hidden_test_access.log"
            access.write_text('{"event":"seal_created"}\n', encoding="utf-8")
            report = export_public_benchmark_manifest(
                ledger_path=root / "ledger.jsonl",
                snapshot=root / "snapshot",
                reviews_path=root / "reviews.jsonl",
                framework_papers_path=root / "framework.jsonl",
                out_dir=root / "out",
                access_log=access,
            )
            self.assertEqual(1, report["paper_count"])
            self.assertEqual(1, report["framework_figure_counts"]["no"])
            overlap = json.loads((root / "out" / "overlap.json").read_text(encoding="utf-8"))
            self.assertEqual("not_downloaded", overlap["sciflow"]["status"])
            self.assertEqual("not_downloaded", overlap["figurebench_official_test"]["status"])
            self.assertIn("public_benchmark_export", access.read_text(encoding="utf-8"))


class ContractDraftTests(unittest.TestCase):
    def test_yes_cites_an_existing_draft_and_unknown_is_queued(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            contracts = root / "structured"
            contracts.mkdir()
            (contracts / "2301.02083.json").write_text("{}", encoding="utf-8")
            (contracts / "2301.02083.validation.json").write_text("{}", encoding="utf-8")
            indexed = index_structured_contracts(contracts)
            self.assertEqual(contracts / "2301.02083.json", indexed["2301.02083"])
            rows = [
                {"paper_id": "yes-paper", "arxiv_id": "2301.02083", "arxiv_ids": ["2301.02083"], "framework_figure": "yes"},
                {"paper_id": "unknown-paper", "arxiv_id": "9999.1", "arxiv_ids": ["9999.1"], "framework_figure": "unknown"},
                {"paper_id": "no-paper", "arxiv_id": "1.2", "arxiv_ids": ["1.2"], "framework_figure": "no"},
                {"paper_id": "train-paper", "arxiv_id": "1.3", "arxiv_ids": ["1.3"], "framework_figure": "yes"},
            ]
            report = draft_contracts(
                rows,
                contracts_dir=contracts,
                out_dir=root / "out",
                exclude_paper_ids={"train-paper"},
            )
            self.assertEqual(0, report["module_edge_n"])
            self.assertFalse(report["signed_confirmation"])
            self.assertEqual(1, report["cited_existing"])
            self.assertEqual(1, report["labeling_queue"])
            self.assertFalse((root / "out" / "confirmation.json").exists())
            drafts = (root / "out" / "contract_drafts.jsonl").read_text(encoding="utf-8")
            self.assertIn("existing_structured_draft", drafts)
            self.assertNotIn("train-paper", drafts)
            self.assertIn("pending", drafts)

    def test_missing_draft_stays_on_the_queue_without_an_extractor(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            report = draft_contracts(
                [{"paper_id": "p", "arxiv_id": "2301.1", "arxiv_ids": ["2301.1"], "framework_figure": "yes"}],
                contracts_dir=root / "missing",
                out_dir=root / "out",
            )
            self.assertEqual(1, report["needs_extraction"])


class ScoreColumnTests(unittest.TestCase):
    def test_flattened_picture_is_not_editable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            picture = root / "full.png"
            Image.new("RGB", (8, 8), "white").save(picture)
            deck = Presentation()
            blank = deck.slide_layouts[6]
            slide = deck.slides.add_slide(blank)
            slide.shapes.add_picture(str(picture), Inches(0), Inches(0), deck.slide_width, deck.slide_height)
            path = root / "flat.pptx"
            deck.save(str(path))
            report = measure_native_editability(path)
            self.assertGreaterEqual(report["raster_area_ratio"], 0.85)
            self.assertEqual("not_editable", report["editability"]["status"])
            self.assertEqual("flattened_picture", report["editability"]["reason"])

    def test_missing_and_non_ooxml_are_not_null(self):
        missing = editability_verdict({"ok": False, "ooxml_valid": False, "error": "missing pptx"})
        self.assertEqual("not_editable", missing["status"])
        raster = editability_verdict({}, delivery="non_ooxml")
        self.assertEqual("non_ooxml", raster["status"])

    def test_paper_quiz_and_nine_stay_null_without_a_legal_response(self):
        self.assertIsNone(parse_paper_quiz({"correct": 1.5, "total": 20}))
        self.assertEqual((18, 20), parse_paper_quiz({"correct": 18, "total": 20}))
        parsed = parse_nine_scores({"aesthetic": 4, "flow": "high", "accuracy": True})
        self.assertEqual(4, parsed["aesthetic"])
        self.assertIsNone(parsed["flow"])
        self.assertIsNone(parsed["accuracy"])
        disabled = score_paper_quiz(None, "paper", spec="off")
        self.assertIsNone(disabled["accuracy"])
        self.assertFalse(disabled["available"])
        nine = score_nine(None, spec="off")
        self.assertIsNone(nine["mean"])

    def test_nine_switches_to_paper_grounded_prompt_when_method_text_is_given(self):
        from rfs.evaluation.livefigure_nine import (
            IMAGE_ONLY_PROMPT_VERSION,
            PAPER_GROUNDED_PROMPT_VERSION,
        )

        image_only = score_nine(None, spec="off")
        self.assertEqual(IMAGE_ONLY_PROMPT_VERSION, image_only["prompt_version"])
        self.assertFalse(image_only["method_grounded"])
        grounded = score_nine(None, method_text="We propose a two-stage encoder.", spec="off")
        self.assertEqual(PAPER_GROUNDED_PROMPT_VERSION, grounded["prompt_version"])
        self.assertTrue(grounded["method_grounded"])
        self.assertNotEqual(image_only["prompt_sha256"], grounded["prompt_sha256"])

        seen: dict = {}

        def fake_call(prompt, images, model):
            seen["prompt"] = prompt
            seen["images"] = images
            return {"aesthetic": 4, "accuracy": 3}

        with tempfile.TemporaryDirectory() as tmp:
            preview = Path(tmp) / "preview.png"
            Image.new("RGB", (4, 4), "white").save(preview)
            report = score_nine(
                preview,
                method_text="We propose a two-stage encoder.",
                spec="api",
                call=fake_call,
            )
        self.assertIn("SOURCE METHOD TEXT", seen["prompt"])
        self.assertIn("two-stage encoder", seen["prompt"])
        self.assertEqual(4, report["aesthetic"])
        self.assertEqual(3, report["accuracy"])
        self.assertEqual(PAPER_GROUNDED_PROMPT_VERSION, report["prompt_version"])

    def test_the_nine_ruler_keeps_its_own_model(self):
        """Scoring reads with a pinned model, so a MODEL_VLM swap cannot move the scores."""
        import os
        from unittest import mock

        from rfs.evaluation.score_model import DEFAULT_SCORE_MODEL, SCORE_MODEL_ENV_VAR, resolve_score_model

        seen: dict = {}

        def fake_call(prompt, images, model):
            seen["model"] = model
            return {"aesthetic": 4}

        with mock.patch.dict(os.environ, {"MODEL_VLM": "gemini-3-pro-preview"}, clear=False):
            os.environ.pop(SCORE_MODEL_ENV_VAR, None)
            self.assertEqual(DEFAULT_SCORE_MODEL, resolve_score_model())
            with tempfile.TemporaryDirectory() as tmp:
                preview = Path(tmp) / "preview.png"
                Image.new("RGB", (4, 4), "white").save(preview)
                report = score_nine(preview, spec="api", call=fake_call)
        self.assertEqual(DEFAULT_SCORE_MODEL, seen["model"])
        self.assertEqual(DEFAULT_SCORE_MODEL, report["model"])
        with mock.patch.dict(os.environ, {SCORE_MODEL_ENV_VAR: "some-other-ruler"}, clear=False):
            self.assertEqual("some-other-ruler", resolve_score_model())

    def test_paper_quiz_reads_with_the_same_pinned_model(self):
        """Both scoring paths name one model, so their columns stay comparable."""
        from rfs.evaluation.score_model import DEFAULT_SCORE_MODEL
        from rfs.evaluation.paper_quiz import write_paper_quiz

        seen: dict = {}

        def fake_call(prompt, images, model):
            seen["model"] = model
            return {"questions": []}

        with tempfile.TemporaryDirectory() as tmp:
            report = write_paper_quiz(
                "We propose a two-stage encoder.",
                Path(tmp) / "questions.json",
                spec="api",
                call=fake_call,
            )
        self.assertEqual(DEFAULT_SCORE_MODEL, seen["model"])
        self.assertFalse(report["ok"])

    def test_clip_is_null_without_method_text_or_an_embedder(self):
        missing = score_paper_figure_clip(Path("absent.png"), "")
        self.assertIsNone(missing["cosine"])
        self.assertEqual("missing method section", missing["reason"])

        class Embedder:
            model_name = "fake-clip"

            def embed_image(self, path):
                return [1.0, 0.0]

            def embed_text(self, text):
                return [1.0, 0.0]

        with tempfile.TemporaryDirectory() as tmp:
            preview = Path(tmp) / "preview.png"
            Image.new("RGB", (4, 4), "white").save(preview)
            scored = score_paper_figure_clip(preview, "method section", embedder=Embedder())
        self.assertEqual(1.0, scored["cosine"])
        self.assertEqual("fake-clip", scored["model"])

    def test_figurebench_splits_are_not_averaged(self):
        high = {
            "aesthetic_quality": 8,
            "expressiveness": 8,
            "polish": 8,
            "clarity": 8,
            "logical_flow": 8,
            "accuracy": 8,
            "completeness": 8,
            "appropriateness": 8,
        }
        low = {key: 2 for key in high}
        record = figurebench_record(official_test=high, development=low)
        self.assertEqual(8.0, record["official_test"]["overall"])
        self.assertEqual(2.0, record["development"]["overall"])
        self.assertIsNone(record["combined"])
        empty = sciflow_record()
        self.assertEqual("official_evaluator_not_connected", empty["reason"])
        self.assertIsNone(empty["node_f1"])

    def test_instance_main_table_omits_the_contract_and_does_not_render(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            instance = root / "instance"
            instance.mkdir()
            (instance / "contract.json").write_text(
                json.dumps(
                    {
                        "expected_figure_content": {
                            "required_entities": ["alpha module"],
                            "required_relations": [{"source": "alpha module", "target": "beta module", "relation_type": "feeds"}],
                        }
                    }
                ),
                encoding="utf-8",
            )
            run = root / "run"
            run.mkdir()
            (run / "figure.tex").write_text("\\node{alpha};", encoding="utf-8")
            metrics = evaluate_instance(instance, run, layers=())
            self.assertNotIn("module", metrics["main_table"])
            self.assertNotIn("edge", metrics["main_table"])
            self.assertEqual("non_ooxml", metrics["main_table"]["editability"]["status"])
            self.assertEqual("disabled", metrics["main_table"]["paper_quiz"]["reason"])
            self.assertIsNone(metrics["main_table"]["preview_render_error"])


class OwnSetScoreTests(unittest.TestCase):
    def test_method_section_stops_before_experiments(self):
        text = "Title\n1 Method\nThe encoder maps tokens.\n2 Experiments\nThe numbers are here.\n"
        section, reason = method_section_from_text(text)
        self.assertIsNone(reason)
        self.assertIn("encoder maps tokens", section)
        self.assertNotIn("numbers are here", section)
        empty, missing = method_section_from_text("Abstract\nNo heading here.\n")
        self.assertEqual("", empty)
        self.assertEqual("missing method section", missing)

    def test_method_heading_tolerates_a_dotted_number(self):
        text = "Title\n3. Method\nOur encoder maps tokens.\n4. Experiments\nThe numbers are here.\n"
        section, reason = method_section_from_text(text)
        self.assertIsNone(reason)
        self.assertIn("encoder maps tokens", section)
        self.assertNotIn("numbers are here", section)
        dotted, dotted_reason = method_section_from_text("2.1. Proposed Method\nA pipeline.\n3 Results\nNumbers.\n")
        self.assertIsNone(dotted_reason)
        self.assertIn("A pipeline.", dotted)
        roman, roman_reason = method_section_from_text("II. Related Work\nEarlier work.\nIII. Method\nOur design.\nIV. Experiments\nNumbers.\n")
        self.assertIsNone(roman_reason)
        self.assertIn("Our design.", roman)
        self.assertNotIn("Earlier work.", roman)
        amp, amp_reason = method_section_from_text("Materials & Methods\nWe culture cells.\nResults\nGrowth.\n")
        self.assertIsNone(amp_reason)
        self.assertIn("We culture cells.", amp)

    def test_question_prompts_ask_about_method_structure(self):
        prompt_dir = Path(__file__).resolve().parents[2] / "benchmarks" / "editable-figure-bench" / "prompts"
        detail = (prompt_dir / "paperquiz-detail-v001.md").read_text(encoding="utf-8")
        understanding = (prompt_dir / "paperquiz-understanding-v001.md").read_text(encoding="utf-8")
        for text in (detail, understanding):
            self.assertIn("Do not ask for the paper title, author names, affiliations", text)
            self.assertNotIn("Title & authorship", text)
        self.assertIn("Named modules", detail)
        self.assertIn("How information moves", understanding)

    def test_na_and_a_missing_reference_are_wrong(self):
        bank = {
            "Question 1": {"aspect": "E", "question": "q", "options": [], "answer": "B. encoder"},
            "Question 2": {"aspect": "E", "question": "q", "options": [], "answer": "A. other"},
        }
        scored = score_answer_bank(
            {
                "Question 1": {"answer": "B", "reference": "center"},
                "Question 2": {"answer": "A", "reference": "NA"},
            },
            bank,
        )
        self.assertEqual(1, scored["correct"])
        self.assertEqual(2, scored["total"])
        self.assertEqual(0.5, scored["accuracy"])
        self.assertEqual(1.0, density_augmented_score(0.5, 100, 100))
        self.assertEqual(0.625, density_augmented_score(0.5, 400, 100))
        self.assertIsNone(density_augmented_score(0.5, 100, 0))

    def test_paper_quiz_asks_from_the_paper_and_grades_the_figure(self):
        paper = "UNIQUE_PAPER_SENTENCE_ZX91 explains the encoder."
        calls = []

        def bank(question: str, answer: str) -> dict:
            return {
                "Question 1": {
                    "aspect": "E",
                    "question": question,
                    "options": ["A. red", "B. encoder", "C. green", "D. gold"],
                    "answer": answer,
                }
            }

        def invoke(prompt, images, model):
            calls.append((prompt, list(images)))
            if not images:
                return bank("What is the module?", "B. encoder") if len(calls) == 1 else bank("What is the goal?", "A. red")
            questions = json.loads(prompt.split("QUESTIONS:\n", 1)[1])
            self.assertNotIn("answer", questions["Question 1"])
            if questions["Question 1"]["question"] == "What is the module?":
                letter = next(option[0] for option in questions["Question 1"]["options"] if "encoder" in option)
                return {"Question 1": {"answer": letter, "reference": "center"}}
            return {"Question 1": {"answer": "NA", "reference": "NA"}}

        with tempfile.TemporaryDirectory() as tmp:
            preview = Path(tmp) / "preview.png"
            Image.new("RGB", (4, 4), "white").save(preview)
            scored = score_paper_quiz(preview, paper, spec="api", call=invoke)
            frozen = json.loads((Path(tmp) / "questions.json").read_text(encoding="utf-8"))
        self.assertEqual(4, len(calls))
        self.assertIn("UNIQUE_PAPER_SENTENCE_ZX91", calls[0][0])
        self.assertIn("UNIQUE_PAPER_SENTENCE_ZX91", calls[1][0])
        self.assertEqual([], calls[0][1])
        self.assertEqual([], calls[1][1])
        self.assertNotIn("UNIQUE_PAPER_SENTENCE_ZX91", calls[2][0])
        self.assertNotIn("UNIQUE_PAPER_SENTENCE_ZX91", calls[3][0])
        self.assertEqual([preview], calls[2][1])
        self.assertIn("answer", frozen["detail"]["Question 1"])
        self.assertEqual(1.0, scored["detail_accuracy"])
        self.assertEqual(0.0, scored["understanding_accuracy"])
        self.assertEqual(0.5, scored["accuracy"])

    def test_frozen_questions_skip_the_paper_call(self):
        calls = []

        def invoke(prompt, images, model):
            calls.append((prompt, list(images)))
            return {"Question 1": {"answer": "B", "reference": "center"}}

        with tempfile.TemporaryDirectory() as tmp:
            preview = Path(tmp) / "preview.png"
            Image.new("RGB", (4, 4), "white").save(preview)
            questions = Path(tmp) / "questions.json"
            questions.write_text(
                json.dumps(
                    {
                        "detail": {
                            "Question 1": {
                                "aspect": "E",
                                "question": "What is the module?",
                                "options": ["A. red", "B. encoder", "C. green", "D. gold"],
                                "answer": "B. encoder",
                            }
                        }
                    }
                ),
                encoding="utf-8",
            )
            scored = score_paper_quiz(
                preview,
                "UNIQUE_PAPER_SENTENCE_ZX91",
                spec="api",
                call=invoke,
                questions_path=questions,
            )
        self.assertEqual(1, len(calls))
        self.assertEqual([preview], calls[0][1])
        self.assertNotIn("UNIQUE_PAPER_SENTENCE_ZX91", calls[0][0])
        self.assertEqual(1.0, scored["accuracy"])
        self.assertEqual(1.0, scored["detail_accuracy"])
        self.assertIsNone(scored["understanding_accuracy"])

    def test_own_delivery_writes_four_columns_and_skips_a_missing_figure(self):
        class Embedder:
            model_name = "fake-clip"

            def embed_image(self, path):
                return [1.0, 0.0]

            def embed_text(self, text):
                return [1.0, 0.0]

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            missing = score_own_delivery(root / "empty", paper_id="missing", paper_quiz_spec="off", nine_spec="off")
            self.assertEqual("delivery_missing", missing["reason"])
            self.assertFalse((root / "empty" / "metrics.json").exists())

            run = root / "10_1038_s41598"
            run.mkdir()
            Image.new("RGB", (8, 8), "white").save(run / "preview.png")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(0.4), Inches(0.4), Inches(3), Inches(0.5)).text_frame.text = "encoder"
            deck.save(str(run / "delivery.pptx"))
            first = score_own_delivery(
                run,
                paper_id="10.1038/s41598",
                method_text="1 Method\nThe encoder maps tokens.\n",
                embedder=Embedder(),
                paper_quiz_spec="off",
                nine_spec="off",
            )
            metrics = json.loads((run / "metrics.json").read_text(encoding="utf-8"))
            second = score_own_delivery(run, paper_id="10.1038/s41598", paper_quiz_spec="off", nine_spec="off")
        self.assertTrue(first["ok"])
        self.assertEqual("v1", metrics["own_table"])
        self.assertFalse(metrics["placeholder"])
        self.assertEqual(1.0, metrics["paper_figure_clip"]["cosine"])
        self.assertEqual("disabled", metrics["paper_quiz"]["reason"])
        self.assertEqual("disabled", metrics["vlm_nine"]["reason"])
        self.assertEqual("editable", metrics["editability"]["status"])
        self.assertEqual("already_scored", second["reason"])

    def test_manifest_scores_only_folders_that_already_have_a_figure(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = root / "manifest.jsonl"
            manifest.write_text('{"paper_id": "doi/a", "pdf_path": "absent.pdf"}\n{"paper_id": "doi/b"}\n', encoding="utf-8")
            delivered = root / "runs" / "doi_b"
            delivered.mkdir(parents=True)
            Image.new("RGB", (4, 4), "white").save(delivered / "figure.png")
            report = score_own_manifest(manifest, root / "runs", paper_quiz_spec="off", nine_spec="off")
            status = json.loads((delivered / "metrics.json").read_text(encoding="utf-8"))["editability"]["status"]
        self.assertEqual(2, report["considered"])
        self.assertEqual(1, report["delivery_missing"])
        self.assertEqual(1, report["scored"])
        self.assertEqual("non_ooxml", status)
        self.assertTrue(report["ok"])


if __name__ == "__main__":
    unittest.main()
