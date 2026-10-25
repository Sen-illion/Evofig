"""Tests for deterministic PowerPoint exemplar object analysis."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from rfs.analysis.reference import analyze_reference_ppt_layout


class ReferencePptLayoutAnalysisTests(unittest.TestCase):
    def test_extracts_picture_provenance_and_rejects_raster_geometry(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image_path = root / "icon.png"
            Image.new("RGB", (120, 80), "white").save(image_path)
            deck = Presentation()
            deck.slide_width = Inches(16)
            deck.slide_height = Inches(9)
            editable = deck.slides.add_slide(deck.slide_layouts[6])
            card = editable.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(5), Inches(3))
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(230, 244, 249)
            text = editable.shapes.add_textbox(Inches(1.4), Inches(2.4), Inches(4), Inches(1))
            run = text.text_frame.paragraphs[0].add_run()
            run.text = "Editable stage"
            run.font.size = Pt(24)
            editable.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6), Inches(3.5), Inches(9), Inches(3.5))
            editable.shapes.add_picture(str(image_path), Inches(10), Inches(2), width=Inches(2), height=Inches(1.3))
            raster = deck.slides.add_slide(deck.slide_layouts[6])
            raster.shapes.add_picture(str(image_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
            deck_path = root / "example.pptx"
            deck.save(deck_path)

            result = analyze_reference_ppt_layout(deck_path, slide_indices=(1, 2), out=root / "analysis")

            self.assertTrue(result["validation"]["ok"], result["validation"]["errors"])
            first, second = result["facts"]["slides"]
            self.assertEqual(first["evidence_mode"], "editable_objects")
            self.assertEqual(first["picture_object_count"], 1)
            self.assertEqual(len(first["picture_assets"][0]["image_sha256"]), 64)
            self.assertEqual(first["typography"]["maximum_pt"], 24.0)
            self.assertEqual(second["evidence_mode"], "raster_only")
            self.assertFalse(second["trusted_for_layout_geometry"])
            self.assertTrue((root / "analysis" / "reference_ppt_layout_facts.json").is_file())


if __name__ == "__main__":
    unittest.main()
