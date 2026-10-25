"""Deterministically inspect editable objects in a PowerPoint exemplar."""

from __future__ import annotations

from collections import Counter
from hashlib import sha256
from pathlib import Path
from statistics import median
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ...contracts import (
    normalize_reference_ppt_layout_facts,
    validate_reference_ppt_layout_facts,
)
from ...utils import ensure_dir, write_json


EMU_PER_INCH = 914400


def _walk_shapes(shapes: Iterable[Any], depth: int = 0, root_shape_id: int | None = None):
    for shape in shapes:
        root_id = int(shape.shape_id) if root_shape_id is None else root_shape_id
        yield shape, depth, root_id
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk_shapes(shape.shapes, depth + 1, root_id)


def _kind(shape: Any) -> str:
    return str(shape.shape_type).split(" (")[0].strip().lower()


def _bbox(shape: Any, slide_width: int, slide_height: int) -> dict[str, float]:
    return {
        "x": round(float(shape.left) / slide_width, 6),
        "y": round(float(shape.top) / slide_height, 6),
        "w": round(float(shape.width) / slide_width, 6),
        "h": round(float(shape.height) / slide_height, 6),
    }


def _text(shape: Any) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return " ".join(str(shape.text_frame.text or "").split())


def _color_hex(color: Any) -> str | None:
    try:
        value = color.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    return f"#{str(value).upper()}" if value is not None else None


def _shape_colors(shape: Any) -> list[str]:
    result: list[str] = []
    try:
        if shape.fill.type == MSO_FILL_TYPE.SOLID:
            if value := _color_hex(shape.fill.fore_color):
                result.append(value)
    except (AttributeError, TypeError, ValueError):
        pass
    try:
        if value := _color_hex(shape.line.color):
            result.append(value)
    except (AttributeError, TypeError, ValueError):
        pass
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    if value := _color_hex(run.font.color):
                        result.append(value)
                except (AttributeError, TypeError, ValueError):
                    pass
    return result


def _font_sizes(shape: Any) -> list[float]:
    if not getattr(shape, "has_text_frame", False):
        return []
    result: list[float] = []
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size is not None:
                result.append(round(float(run.font.size.pt), 2))
    return result


def _union_area_ratio(boxes: list[dict[str, float]], resolution: int = 120) -> float:
    rows = max(1, round(resolution * 9 / 16))
    occupied: set[tuple[int, int]] = set()
    for box in boxes:
        x0 = max(0, min(resolution, int(box["x"] * resolution)))
        x1 = max(x0, min(resolution, round((box["x"] + box["w"]) * resolution)))
        y0 = max(0, min(rows, int(box["y"] * rows)))
        y1 = max(y0, min(rows, round((box["y"] + box["h"]) * rows)))
        occupied.update((x, y) for x in range(x0, x1) for y in range(y0, y1))
    return round(len(occupied) / max(1, resolution * rows), 4)


def _picture_asset(shape: Any, root_shape_id: int, slide_width: int, slide_height: int) -> dict[str, Any]:
    blob = bytes(shape.image.blob)
    width, height = shape.image.size
    return {
        "shape_id": int(shape.shape_id),
        "name": str(shape.name),
        "containing_group_shape_id": root_shape_id if root_shape_id != int(shape.shape_id) else 0,
        "bbox_percent": _bbox(shape, slide_width, slide_height),
        "image_sha256": sha256(blob).hexdigest(),
        "image_extension": str(shape.image.ext or "").lower(),
        "pixel_width": int(width),
        "pixel_height": int(height),
    }


def _slide_facts(slide: Any, index: int, slide_width: int, slide_height: int) -> dict[str, Any]:
    roots = list(slide.shapes)
    walked = list(_walk_shapes(roots))
    root_descendants: dict[int, Counter[str]] = {}
    for shape, _depth, root_id in walked:
        root_descendants.setdefault(root_id, Counter())[_kind(shape)] += 1
    pictures = [
        _picture_asset(shape, root_id, slide_width, slide_height)
        for shape, _depth, root_id in walked
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    full_slide_picture = (
        len(roots) == 1
        and roots[0].shape_type == MSO_SHAPE_TYPE.PICTURE
        and float(roots[0].width) * float(roots[0].height) / (slide_width * slide_height) >= 0.85
    )
    font_sizes = [size for shape, _depth, _root in walked for size in _font_sizes(shape)]
    colors = Counter(color for shape, _depth, _root in walked for color in _shape_colors(shape))
    object_counts = Counter(_kind(shape) for shape, _depth, _root in walked)
    native_count = sum(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape, _depth, _root in walked)
    return {
        "slide_index": index,
        "evidence_mode": "raster_only" if full_slide_picture else "editable_objects",
        "trusted_for_layout_geometry": not full_slide_picture,
        "root_shape_count": len(roots),
        "recursive_object_count": len(walked),
        "maximum_group_depth": max((depth for _shape, depth, _root in walked), default=0),
        "object_counts": dict(sorted(object_counts.items())),
        "editable_object_ratio": round(native_count / max(1, len(walked)), 4),
        "native_editable_object_count": native_count,
        "picture_object_count": len(pictures),
        "root_region_occupancy": _union_area_ratio([_bbox(shape, slide_width, slide_height) for shape in roots]),
        "root_regions": [
            {
                "shape_id": int(shape.shape_id),
                "name": str(shape.name),
                "kind": _kind(shape),
                "bbox_percent": _bbox(shape, slide_width, slide_height),
                "text_preview": _text(shape)[:160],
                "descendant_counts": dict(sorted(root_descendants[int(shape.shape_id)].items())),
            }
            for shape in roots
        ],
        "picture_assets": pictures,
        "typography": {
            "explicit_run_count": len(font_sizes),
            "minimum_pt": min(font_sizes) if font_sizes else None,
            "median_pt": round(median(font_sizes), 2) if font_sizes else None,
            "maximum_pt": max(font_sizes) if font_sizes else None,
        },
        "object_color_frequency": [
            {"hex": color, "object_occurrences": count}
            for color, count in colors.most_common(16)
        ],
        "learning_policy": (
            "visual reference only; editable hierarchy and connector geometry are not evidenced"
            if full_slide_picture
            else "deterministic editable-object facts; transferable judgments require human review"
        ),
    }


def analyze_reference_ppt_layout(
    pptx: str | Path,
    *,
    slide_indices: Iterable[int] = (1,),
    out: str | Path | None = None,
) -> dict[str, Any]:
    """Inspect selected slides and optionally write facts plus validation reports."""

    source = Path(pptx).resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.suffix.lower() != ".pptx":
        raise ValueError(f"Expected PPTX input: {source}")
    presentation = Presentation(source)
    selected = sorted(set(int(index) for index in slide_indices))
    if not selected or selected[0] < 1 or selected[-1] > len(presentation.slides):
        raise ValueError(f"slide_indices must be within 1..{len(presentation.slides)}")
    raw = {
        "summary": "Deterministic PPTX exemplar object facts for selected slides.",
        "source": {
            "path": str(source),
            "sha256": sha256(source.read_bytes()).hexdigest(),
            "total_slide_count": len(presentation.slides),
        },
        "canvas": {
            "width_inches": round(presentation.slide_width / EMU_PER_INCH, 4),
            "height_inches": round(presentation.slide_height / EMU_PER_INCH, 4),
            "aspect_ratio": round(presentation.slide_width / presentation.slide_height, 6),
        },
        "slides": [
            _slide_facts(presentation.slides[index - 1], index, presentation.slide_width, presentation.slide_height)
            for index in selected
        ],
        "safety": {
            "scientific_semantics_learned": False,
            "exact_text_reusable": False,
            "raster_only_geometry_reusable": False,
        },
    }
    facts = normalize_reference_ppt_layout_facts(raw)
    validation = validate_reference_ppt_layout_facts(facts)
    if out is not None:
        output = ensure_dir(out)
        write_json(output / "reference_ppt_layout_facts.json", facts)
        write_json(output / "reference_ppt_layout_validation.json", {
            "summary": "Validation for deterministic PPTX exemplar object facts.",
            **{key: value for key, value in validation.items() if key != "facts"},
        })
    return {"facts": facts, "validation": {key: value for key, value in validation.items() if key != "facts"}}


__all__ = ["analyze_reference_ppt_layout"]
